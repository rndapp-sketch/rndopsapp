"""Loads the externally-generated migration mapping into Legacy Project Mapping.

Three workbooks feed this, all produced by the scripts in the
`merging_excel_2_pi_on_pragati_empInitials` toolset:

  Pragati_Migration_Mapping.xlsx        old project <-> Pragati project (sheet "Migration Mapping")
  PROJECT_Excel_Info.xlsx               Proman file index            (sheet "Extracted Files")
  RNDOPS_PROJECTS_REMOTE_Excel_Info.xlsx  R&D OPS file index         (sheet "Extracted Files")

Re-running is safe: each Pragati project's mapping row is replaced wholesale,
so a regenerated mapping simply overwrites the previous one.
"""

import os
import re
from datetime import datetime

import frappe
import openpyxl

MAPPING_SHEET = "Migration Mapping"
INDEX_SHEET = "Extracted Files"
COMMIT_EVERY = 50

SOURCE_SYSTEMS = ("Proman", "R&D OPS")

# Legacy filenames carry noise around the project code -- "-reconciled",
# " Final UC", surrounding parentheses -- so the code is extracted rather than
# compared whole. Format: 4-char dept, 3-char type, 4-char agency, 5-digit
# employee id, 3-4 char initials, 3-digit serial.
CODE_RE = re.compile(r"[A-Z0-9]{4}[A-Z]{3}[A-Z0-9]{4}\d{5}[A-Z0-9]{3,4}\d{3}")


def _normalize_code(value):
	return re.sub(r"[^A-Za-z0-9]", "", str(value or "")).upper()


def _extract_code(value):
	match = CODE_RE.search(_normalize_code(value))
	return match.group(0) if match else ""


def _rel_path(folder_path, file_name):
	"""Path relative to the source system's own PROJECT root, forward-slashed.

	Both index workbooks store absolute paths from the machine that scanned
	them, but every one of them contains the system's PROJECT root as a
	segment -- which is also the directory Tomcat serves in production.
	"""
	folder = str(folder_path or "").replace("\\", "/")
	if "/PROJECT/" not in folder:
		return ""
	return folder.rsplit("/PROJECT/", 1)[1].strip("/") + "/" + str(file_name)


def _read_sheet(path, sheet_name):
	"""Yield each row of a sheet as a dict keyed by its header text."""
	if not os.path.isfile(path):
		frappe.throw(f"Workbook not found: {path}")

	book = openpyxl.load_workbook(path, read_only=True, data_only=True)
	try:
		if sheet_name not in book.sheetnames:
			frappe.throw(f"Sheet '{sheet_name}' not found in {os.path.basename(path)}")
		sheet = book[sheet_name]
		rows = sheet.iter_rows(values_only=True)
		headers = [str(h).strip() if h is not None else "" for h in next(rows, [])]
		for row in rows:
			yield dict(zip(headers, row))
	finally:
		book.close()


def build_file_index(proman_index_path, rndops_index_path):
	"""(source_system, project_code) -> [{file_name, rel_path, file_ext}]"""
	index = {}

	for source_system, path in (("Proman", proman_index_path), ("R&D OPS", rndops_index_path)):
		for row in _read_sheet(path, INDEX_SHEET):
			file_name = row.get("Excel File Name")
			code = _extract_code(row.get("Project Number") or file_name)
			rel_path = _rel_path(row.get("Folder Path"), file_name)
			if not (code and rel_path):
				continue
			index.setdefault((source_system, code), []).append(
				{
					"file_name": str(file_name),
					"rel_path": rel_path,
					"file_ext": os.path.splitext(str(file_name))[1].lower(),
				}
			)

	for files in index.values():
		files.sort(key=lambda f: f["file_name"])

	return index


def _upsert(docname, mapping_row, files, batch):
	if frappe.db.exists("Legacy Project Mapping", docname):
		doc = frappe.get_doc("Legacy Project Mapping", docname)
	else:
		doc = frappe.new_doc("Legacy Project Mapping")
		doc.project_registration = docname

	doc.source_system = mapping_row["source_system"]
	doc.old_project_number = mapping_row["old_project_number"]
	doc.old_project_title = mapping_row["old_project_title"]
	doc.old_pi_name = mapping_row["old_pi_name"]
	doc.old_status = mapping_row["old_status"]
	doc.match_method = mapping_row["match_method"]
	doc.title_similarity = mapping_row["title_similarity"]
	doc.import_batch = batch
	doc.last_imported_on = datetime.now()

	doc.set("files", [])
	for entry in files:
		doc.append("files", entry)

	doc.flags.ignore_permissions = True
	doc.save()
	return doc


def _float_or_none(value):
	try:
		return float(value)
	except (TypeError, ValueError):
		return None


def import_legacy_mapping(mapping_path, proman_index_path, rndops_index_path, dry_run=False):
	"""Upsert one Legacy Project Mapping per migrated Pragati project."""
	batch = datetime.now().strftime("%Y%m%d%H%M%S")
	file_index = build_file_index(proman_index_path, rndops_index_path)

	stats = {
		"batch": batch,
		"dry_run": bool(dry_run),
		"indexed_files": sum(len(v) for v in file_index.values()),
		"migrated_rows": 0,
		"created": 0,
		"updated": 0,
		"with_files": 0,
		"without_files": 0,
		"unknown_project_no": 0,
		"removed_stale": 0,
		"warnings": [],
	}
	touched = 0

	for row in _read_sheet(mapping_path, MAPPING_SHEET):
		if str(row.get("Migrated to Pragati") or "").strip().lower() != "yes":
			continue
		stats["migrated_rows"] += 1

		pragati_no = str(row.get("Pragati Project No") or "").strip()
		source_system = str(row.get("Source System") or "").strip()
		if not pragati_no or source_system not in SOURCE_SYSTEMS:
			stats["warnings"].append(f"Skipped row with project_no={pragati_no!r} source={source_system!r}")
			continue

		# project_no is not unique in Pragati -- a handful are shared by several
		# Project Registration docs, and each of them gets its own mapping.
		docnames = frappe.get_all(
			"Project Registration", filters={"project_no": pragati_no}, pluck="name"
		)
		if not docnames:
			stats["unknown_project_no"] += 1
			stats["warnings"].append(f"No Project Registration with project_no {pragati_no}")
			continue

		old_number = str(row.get("Old Project Number") or "").strip()
		files = file_index.get((source_system, _normalize_code(old_number)), [])
		if files:
			stats["with_files"] += 1
		else:
			stats["without_files"] += 1

		mapping_row = {
			"source_system": source_system,
			"old_project_number": old_number,
			"old_project_title": row.get("Old Project Title"),
			"old_pi_name": row.get("Old PI Name"),
			"old_status": row.get("Old Status"),
			"match_method": row.get("Match Method"),
			"title_similarity": _float_or_none(row.get("Title Similarity %")),
		}

		for docname in docnames:
			existed = frappe.db.exists("Legacy Project Mapping", docname)
			if not dry_run:
				_upsert(docname, mapping_row, files, batch)
				touched += 1
				if touched % COMMIT_EVERY == 0:
					frappe.db.commit()
			stats["updated" if existed else "created"] += 1

	# A regenerated mapping can drop a project that previously matched. Anything
	# this run did not touch is no longer a confirmed migration, so it goes --
	# otherwise the project would keep showing a legacy record that the matcher
	# has since rejected.
	# (Skipped on a dry run, where nothing was stamped with this batch and every
	# existing record would therefore look stale.)
	# `touched` guards against a wrong or empty workbook pruning every mapping.
	if not dry_run and touched:
		stale = frappe.get_all(
			"Legacy Project Mapping", filters={"import_batch": ["!=", batch]}, pluck="name"
		)
		stats["removed_stale"] = len(stale)
		for docname in stale:
			frappe.delete_doc("Legacy Project Mapping", docname, ignore_permissions=True, force=True)
		frappe.db.commit()

	return stats
